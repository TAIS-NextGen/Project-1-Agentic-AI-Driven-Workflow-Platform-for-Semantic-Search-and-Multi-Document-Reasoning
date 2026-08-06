from __future__ import annotations

from pathlib import Path

import pytest

from backend.nodes.ingestion.document_parser import DocumentParserNode
from backend.sdk import ExecutionContext, NodeStatus


def _context(path: Path) -> ExecutionContext:
    return ExecutionContext(
        workflow_id="wf-parser",
        node_id="parser-1",
        inputs={
            "document": {
                "filename": path.name,
                "path": str(path),
                "size_bytes": path.stat().st_size,
            }
        },
        config={},
        services={},
    )


def test_node_definition():
    node = DocumentParserNode(node_id="parser-1", config={})
    assert node.type == "document-parser"
    assert node.inputs[0].name == "document"
    assert {port.name for port in node.outputs} == {"text", "data", "metadata", "downloads"}


@pytest.mark.asyncio
async def test_parses_docx_text_and_tables(tmp_path: Path):
    import docx

    path = tmp_path / "report.docx"
    document = docx.Document()
    document.add_heading("Quarterly Report", level=1)
    document.add_paragraph("Revenue increased by 12 percent.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "120"
    document.save(path)

    node = DocumentParserNode(node_id="parser-1", config={"include_tables": True})
    result = await node.execute(_context(path))

    assert result.status == NodeStatus.SUCCESS, result.error
    assert "Quarterly Report" in result.outputs["text"]
    assert "Revenue\t120" in result.outputs["text"]
    assert result.outputs["metadata"]["parser_engine"] == "python-docx"
    assert result.outputs["metadata"]["table_count"] == 1


@pytest.mark.asyncio
async def test_parses_xlsx_as_text_and_structured_rows(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "sales.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Product", "Amount"])
    sheet.append(["A", 10])
    sheet.append(["B", 20])
    workbook.save(path)

    node = DocumentParserNode(node_id="parser-1", config={})
    result = await node.execute(_context(path))

    assert result.status == NodeStatus.SUCCESS, result.error
    assert "[Sales]" in result.outputs["text"]
    assert "Product\tAmount" in result.outputs["text"]
    sheets = result.outputs["data"]["sheets"]
    assert sheets[0]["name"] == "Sales"
    assert sheets[0]["row_count"] == 3
    assert sheets[0]["rows"][1] == ["A", 10]
    assert result.outputs["metadata"]["parser_engine"] == "openpyxl"


@pytest.mark.asyncio
async def test_parses_pdf_pages(tmp_path: Path):
    from pypdf import PdfWriter

    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as stream:
        writer.write(stream)

    node = DocumentParserNode(node_id="parser-1", config={})
    result = await node.execute(_context(path))

    assert result.status == NodeStatus.SUCCESS, result.error
    assert result.outputs["metadata"]["parser_engine"] == "pypdf"
    assert result.outputs["metadata"]["page_count"] == 1
    assert result.outputs["data"]["pages"][0]["page"] == 1


@pytest.mark.asyncio
async def test_rejects_legacy_doc_format(tmp_path: Path):
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy")
    node = DocumentParserNode(node_id="parser-1", config={})

    result = await node.execute(_context(path))

    assert result.status == NodeStatus.FAILURE
    assert "legacy .doc/.xls files must be converted" in (result.error or "")
