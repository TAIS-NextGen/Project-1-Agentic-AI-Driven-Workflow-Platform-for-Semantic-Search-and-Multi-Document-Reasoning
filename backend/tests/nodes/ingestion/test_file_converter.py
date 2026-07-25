from __future__ import annotations

from pathlib import Path

import pytest

from backend.nodes.ingestion.file_converter import FileConverterNode
from backend.sdk import ExecutionContext, NodeStatus


@pytest.fixture
def node():
    return FileConverterNode(node_id="test-converter-1", config={})


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    import docx

    f = tmp_path / "test.docx"
    document = docx.Document()
    document.add_paragraph("Hello, world!")
    document.add_paragraph("Second paragraph.")
    document.save(str(f))
    return f


@pytest.fixture
def sample_png(tmp_path: Path) -> Path:
    from PIL import Image

    f = tmp_path / "test.png"
    image = Image.new("RGB", (100, 100), color="white")
    image.save(f)
    return f


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    f = tmp_path / "test.pdf"
    f.write_bytes(b"%PDF-1.4 mock pdf content")
    return f


@pytest.fixture
def upload_dir(monkeypatch, tmp_path: Path) -> Path:
    d = tmp_path / "uploads"
    d.mkdir(parents=True, exist_ok=True)
    monkeypatch.setenv("UPLOAD_DIR", str(d))
    return d


def test_node_definition(node: FileConverterNode):
    assert node.type == "file-converter"
    assert node.name == "File Converter"
    assert node.category == "ingestion"
    assert len(node.inputs) == 1
    assert node.inputs[0].name == "document"
    assert len(node.outputs) == 2
    output_names = {p.name for p in node.outputs}
    assert output_names == {"converted_path", "text"}
    assert len(node.config_fields) == 2


@pytest.mark.asyncio
async def test_execute_pdf_passthrough(node: FileConverterNode, sample_pdf: Path):
    ctx = ExecutionContext(
        workflow_id="wf-1",
        node_id="test-converter-1",
        inputs={
            "document": {
                "filename": sample_pdf.name,
                "path": str(sample_pdf),
                "size_bytes": sample_pdf.stat().st_size,
            }
        },
        config={},
        services={},
    )

    result = await node.execute(ctx)

    assert result.status == NodeStatus.SUCCESS
    assert result.outputs["converted_path"] == str(sample_pdf)


@pytest.mark.asyncio
async def test_execute_image_to_pdf(node: FileConverterNode, sample_png: Path):
    ctx = ExecutionContext(
        workflow_id="wf-1",
        node_id="test-converter-1",
        inputs={
            "document": {
                "filename": sample_png.name,
                "path": str(sample_png),
                "size_bytes": sample_png.stat().st_size,
            }
        },
        config={},
        services={},
    )

    result = await node.execute(ctx)

    assert result.status == NodeStatus.SUCCESS
    output_path = Path(result.outputs["converted_path"])
    assert output_path.exists()
    assert output_path.suffix == ".pdf"


@pytest.mark.asyncio
async def test_execute_docx_text_extraction(node: FileConverterNode, sample_docx: Path):
    ctx = ExecutionContext(
        workflow_id="wf-1",
        node_id="test-converter-1",
        inputs={
            "document": {
                "filename": sample_docx.name,
                "path": str(sample_docx),
                "size_bytes": sample_docx.stat().st_size,
            }
        },
        config={"output_mode": "text"},
        services={},
    )

    result = await node.execute(ctx)

    assert result.status == NodeStatus.SUCCESS
    assert "Hello, world!" in result.outputs["text"]
    assert "Second paragraph." in result.outputs["text"]


@pytest.mark.asyncio
async def test_execute_docx_to_pdf(sample_docx: Path):
    node = FileConverterNode(node_id="test-converter-1", config={"output_mode": "pdf"})

    ctx = ExecutionContext(
        workflow_id="wf-1",
        node_id="test-converter-1",
        inputs={
            "document": {
                "filename": sample_docx.name,
                "path": str(sample_docx),
                "size_bytes": sample_docx.stat().st_size,
            }
        },
        config={"output_mode": "pdf"},
        services={},
    )

    result = await node.execute(ctx)

    assert result.status == NodeStatus.SUCCESS
    output_path = Path(result.outputs["converted_path"])
    assert output_path.exists()
    assert output_path.suffix == ".pdf"


@pytest.mark.asyncio
async def test_execute_fails_with_no_input(node: FileConverterNode):
    ctx = ExecutionContext(
        workflow_id="wf-1",
        node_id="test-converter-1",
        inputs={},
        config={},
        services={},
    )

    result = await node.execute(ctx)

    assert result.status == NodeStatus.FAILURE
    assert "No document provided" in result.error


@pytest.mark.asyncio
async def test_execute_docx_text_extraction(sample_docx: Path):
    node = FileConverterNode(node_id="test-converter-1", config={"output_mode": "text"})

    ctx = ExecutionContext(
        workflow_id="wf-1",
        node_id="test-converter-1",
        inputs={
            "document": {
                "filename": sample_docx.name,
                "path": str(sample_docx),
                "size_bytes": sample_docx.stat().st_size,
            }
        },
        config={"output_mode": "text"},
        services={},
    )

    result = await node.execute(ctx)

    assert result.status == NodeStatus.SUCCESS
    assert "Hello, world!" in result.outputs["text"]
    assert "Second paragraph." in result.outputs["text"]


def test_to_definition_returns_correct_shape(node: FileConverterNode):
    definition = node.to_definition()
    assert definition["type"] == "file-converter"
    assert definition["name"] == "File Converter"
    assert definition["category"] == "ingestion"
    assert len(definition["inputs"]) == 1
    assert len(definition["outputs"]) == 2
    assert len(definition["config_fields"]) == 2