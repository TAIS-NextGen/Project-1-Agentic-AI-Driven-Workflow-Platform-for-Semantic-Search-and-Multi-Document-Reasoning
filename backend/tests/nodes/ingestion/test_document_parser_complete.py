from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from backend.nodes.ingestion.document_parser import DocumentParserNode
from backend.sdk import ExecutionContext, NodeStatus


def make_context(path: Path, config: dict | None = None, inputs: dict | None = None) -> ExecutionContext:
    resolved_inputs = inputs or {
        "document": {
            "filename": path.name,
            "path": str(path),
            "size_bytes": path.stat().st_size,
        }
    }
    return ExecutionContext(
        workflow_id="wf-parser-tests",
        node_id="parser-1",
        inputs=resolved_inputs,
        config=config or {},
    )


@pytest.mark.asyncio
async def test_parser_selects_nearest_processed_upstream_artifact(tmp_path: Path):
    original = tmp_path / "original.txt"
    processed = tmp_path / "processed.txt"
    original.write_text("ORIGINAL VERSION", encoding="utf-8")
    processed.write_text("LATEST PROCESSED VERSION", encoding="utf-8")

    inputs = {
        "__upstream_artifacts__": [
            {
                "source_node_id": "cleaner-1",
                "source_node_type": "file-converter",
                "output_port": "converted_path",
                "distance": 1,
                "value": str(processed),
            },
            {
                "source_node_id": "upload-1",
                "source_node_type": "document-upload",
                "output_port": "document",
                "distance": 2,
                "value": {"path": str(original), "filename": original.name},
            },
        ]
    }
    config = {"parse_entire_document": True, "create_downloads": False}
    node = DocumentParserNode(node_id="parser-1", config=config)
    result = await node.execute(make_context(processed, config, inputs))

    assert result.status == NodeStatus.SUCCESS, result.error
    assert result.outputs["text"] == "LATEST PROCESSED VERSION"
    assert result.outputs["metadata"]["source_node_id"] == "cleaner-1"
    assert result.outputs["metadata"]["source_output_port"] == "converted_path"
    assert result.outputs["metadata"]["selected_source"] == "nearest_upstream_artifact"


@pytest.mark.asyncio
async def test_parse_entire_excel_ignores_row_limit(tmp_path: Path):
    path = tmp_path / "large.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("All rows")
    for index in range(1, 5006):
        sheet.append([index, f"row-{index}"])
    workbook.save(path)

    config = {
        "parse_entire_document": True,
        "include_tables": True,
        "max_rows_per_sheet": 5,
        "create_downloads": False,
    }
    node = DocumentParserNode(node_id="parser-1", config=config)
    result = await node.execute(make_context(path, config))

    assert result.status == NodeStatus.SUCCESS, result.error
    assert result.outputs["data"]["sheets"][0]["row_count"] == 5005
    assert result.outputs["data"]["sheets"][0]["truncated"] is False
    assert result.outputs["metadata"]["complete"] is True
    assert "row-5005" in result.outputs["text"]


@pytest.mark.asyncio
async def test_limited_excel_reports_truncation(tmp_path: Path):
    path = tmp_path / "limited.xlsx"
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Limited")
    for index in range(1, 21):
        sheet.append([index, f"row-{index}"])
    workbook.save(path)

    config = {
        "parse_entire_document": False,
        "include_tables": True,
        "max_rows_per_sheet": 5,
        "create_downloads": False,
    }
    node = DocumentParserNode(node_id="parser-1", config=config)
    result = await node.execute(make_context(path, config))

    assert result.status == NodeStatus.SUCCESS, result.error
    assert result.outputs["data"]["sheets"][0]["row_count"] == 5
    assert result.outputs["data"]["sheets"][0]["truncated"] is True
    assert result.outputs["metadata"]["complete"] is False
    assert result.outputs["metadata"]["truncated"] is True


@pytest.mark.asyncio
async def test_docx_full_parse_includes_ordered_blocks_headers_and_footers(tmp_path: Path):
    path = tmp_path / "complete.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "HEADER CONTENT"
    document.add_paragraph("FIRST PARAGRAPH")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    document.add_paragraph("LAST PARAGRAPH")
    document.sections[0].footer.paragraphs[0].text = "FOOTER CONTENT"
    document.save(path)

    config = {
        "parse_entire_document": True,
        "include_tables": True,
        "include_headers_footers": True,
        "create_downloads": False,
    }
    node = DocumentParserNode(node_id="parser-1", config=config)
    result = await node.execute(make_context(path, config))

    assert result.status == NodeStatus.SUCCESS, result.error
    text = result.outputs["text"]
    assert "HEADER CONTENT" in text
    assert "FIRST PARAGRAPH" in text
    assert "A1\tB1" in text
    assert "LAST PARAGRAPH" in text
    assert "FOOTER CONTENT" in text
    block_types = [block["type"] for block in result.outputs["data"]["blocks"]]
    assert block_types == ["paragraph", "table", "paragraph"]
    assert result.outputs["metadata"]["complete"] is True


@pytest.mark.asyncio
async def test_parser_creates_complete_download_files(tmp_path: Path, monkeypatch):
    source = tmp_path / "rapport.txt"
    source.write_text("ligne 1\nligne 2\nمرحبا", encoding="utf-8")
    storage = tmp_path / "storage"
    monkeypatch.setenv("STORAGE_DIR", str(storage))

    config = {"parse_entire_document": True, "create_downloads": True}
    node = DocumentParserNode(node_id="parser-1", config=config)
    result = await node.execute(make_context(source, config))

    assert result.status == NodeStatus.SUCCESS, result.error
    downloads = result.outputs["downloads"]
    text_path = Path(downloads["text"]["path"])
    json_path = Path(downloads["json"]["path"])
    bundle_path = Path(downloads["bundle"]["path"])
    assert text_path.exists()
    assert json_path.exists()
    assert bundle_path.exists()
    assert text_path.read_text(encoding="utf-8") == result.outputs["text"]
    assert downloads["text"]["download_url"].startswith("/data/storage/parser_outputs/")

    with zipfile.ZipFile(bundle_path) as archive:
        names = archive.namelist()
        assert downloads["text"]["filename"] in names
        assert downloads["json"]["filename"] in names
