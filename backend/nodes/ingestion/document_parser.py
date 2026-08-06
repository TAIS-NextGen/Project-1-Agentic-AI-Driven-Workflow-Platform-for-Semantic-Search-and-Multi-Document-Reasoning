from __future__ import annotations

import csv
import io
import json
import mimetypes
import os
import re
import uuid
import zipfile
from pathlib import Path
from typing import Any, Iterable

from backend.sdk import (
    BaseNode,
    ConfigField,
    ExecutionContext,
    NodeResult,
    Port,
    PortType,
)


class DocumentParserNode(BaseNode):
    type = "document-parser"
    name = "Document Parser"
    category = "ingestion"
    icon = "🧩"
    color = "#14b8a6"
    description = "Extract the complete text and structured data from PDF, Word, Excel, CSV, and text documents"
    version = "1.1.0"

    inputs = [
        Port(
            name="document",
            type=PortType.DOCUMENT,
            label="Latest Document",
            description="The latest document artifact produced by the connected upstream chain",
            required=False,
        ),
    ]

    outputs = [
        Port(
            name="text",
            type=PortType.TEXT,
            label="Full Raw Text",
            description="Complete extracted textual content",
        ),
        Port(
            name="data",
            type=PortType.JSON,
            label="Structured Data",
            description="Pages, ordered Word blocks, tables, sheets, or rows when available",
        ),
        Port(
            name="metadata",
            type=PortType.JSON,
            label="Parser Metadata",
            description="Document type, extraction statistics, source lineage, and truncation status",
        ),
        Port(
            name="downloads",
            type=PortType.JSON,
            label="Downloadable Outputs",
            description="Download links for full TXT, structured JSON, and ZIP bundle outputs",
        ),
    ]

    config_fields = [
        ConfigField(
            key="file_id",
            label="File ID",
            type="text",
            required=False,
            default="",
            description="Standalone fallback only. A connected upstream document always takes priority.",
        ),
        ConfigField(
            key="parse_entire_document",
            label="Parse Entire Document",
            type="boolean",
            required=False,
            default=True,
            description="Read every PDF page, Word block, Excel sheet/row, CSV row, and text line.",
        ),
        ConfigField(
            key="include_tables",
            label="Include Tables",
            type="boolean",
            required=False,
            default=True,
            description="Include Word tables and spreadsheet rows in structured output.",
        ),
        ConfigField(
            key="include_headers_footers",
            label="Include Headers and Footers",
            type="boolean",
            required=False,
            default=True,
            description="Extract Word section headers and footers as part of the complete document.",
        ),
        ConfigField(
            key="preserve_page_breaks",
            label="Preserve PDF Page Breaks",
            type="boolean",
            required=False,
            default=True,
            description="Insert visible page separators in the full text output.",
        ),
        ConfigField(
            key="sheet_name",
            label="Excel Sheet Name",
            type="text",
            required=False,
            default="",
            description="Optional Excel sheet to parse. Empty means all sheets.",
        ),
        ConfigField(
            key="max_rows_per_sheet",
            label="Max Rows per Sheet",
            type="number",
            required=False,
            default=5000,
            description="Used only when 'Parse Entire Document' is disabled.",
        ),
        ConfigField(
            key="max_file_size_mb",
            label="Max File Size (MB)",
            type="number",
            required=False,
            default=50,
        ),
        ConfigField(
            key="create_downloads",
            label="Create Downloadable Outputs",
            type="boolean",
            required=False,
            default=True,
            description="Create TXT, JSON, and ZIP files containing the complete parser result.",
        ),
    ]

    _PATH_KEYS = (
        "path",
        "file_path",
        "output_path",
        "converted_path",
        "cleaned_document",
        "converted_file",
    )
    _NESTED_ARTIFACT_KEYS = (
        "document",
        "image",
        "file",
        "artifact",
        "latest_document",
        "converted_document",
        "cleaned_document",
        "converted_path",
        "output_path",
    )

    @staticmethod
    def _get_upload_dir() -> Path:
        return Path(os.getenv("UPLOAD_DIR") or os.getenv("SYMPACT_UPLOAD_DIR") or "data/uploads")

    @staticmethod
    def _get_storage_dir() -> Path:
        return Path(os.getenv("STORAGE_DIR") or os.getenv("SYMPACT_STORAGE_DIR") or "data/storage")

    @classmethod
    def _candidate_to_file_data(
        cls,
        candidate: Any,
        *,
        provenance: dict[str, Any] | None = None,
        visited: set[int] | None = None,
    ) -> dict[str, Any] | None:
        """Normalize strings and nested node outputs into one document artifact.

        Upstream processors in this project do not all use the same output name. Some
        return ``document`` dictionaries, while older nodes return a path string under
        names such as ``converted_path`` or ``cleaned_document``. This method accepts
        all of those representations without forcing every existing node to change.
        """
        if candidate is None:
            return None

        visited = visited or set()
        marker = id(candidate)
        if marker in visited:
            return None
        visited.add(marker)

        if isinstance(candidate, Path):
            candidate = str(candidate)

        if isinstance(candidate, str):
            value = candidate.strip()
            if not value:
                return None
            path = Path(value)
            if not path.exists() or not path.is_file():
                return None
            payload: dict[str, Any] = {
                "filename": path.name,
                "path": str(path),
                "file_path": str(path),
                "size_bytes": path.stat().st_size,
                "mime_type": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
            }
            if provenance:
                payload["_provenance"] = provenance
            return payload

        if isinstance(candidate, dict):
            for key in cls._PATH_KEYS:
                raw_path = candidate.get(key)
                if isinstance(raw_path, (str, Path)) and str(raw_path).strip():
                    path = Path(str(raw_path))
                    if path.exists() and path.is_file():
                        payload = dict(candidate)
                        payload["path"] = str(path)
                        payload["file_path"] = str(path)
                        payload.setdefault("filename", path.name)
                        payload.setdefault("size_bytes", path.stat().st_size)
                        payload.setdefault(
                            "mime_type",
                            mimetypes.guess_type(str(payload.get("filename") or path.name))[0]
                            or "application/octet-stream",
                        )
                        if provenance:
                            payload["_provenance"] = provenance
                        return payload

            for key in cls._NESTED_ARTIFACT_KEYS:
                if key not in candidate:
                    continue
                nested = cls._candidate_to_file_data(
                    candidate[key],
                    provenance=provenance,
                    visited=visited,
                )
                if nested:
                    for metadata_key in ("filename", "mime_type", "size_bytes", "file_id"):
                        if metadata_key in candidate and metadata_key not in nested:
                            nested[metadata_key] = candidate[metadata_key]
                    return nested

            # Finally search arbitrary nested values. This is intentionally last so a
            # nearby transformed document is selected before unrelated metadata paths.
            for value in candidate.values():
                if isinstance(value, (dict, list, tuple, str, Path)):
                    nested = cls._candidate_to_file_data(
                        value,
                        provenance=provenance,
                        visited=visited,
                    )
                    if nested:
                        return nested

        if isinstance(candidate, (list, tuple)):
            # The executor supplies lineage nearest-first; preserve that ordering.
            for value in candidate:
                nested = cls._candidate_to_file_data(
                    value,
                    provenance=provenance,
                    visited=visited,
                )
                if nested:
                    return nested

        return None

    @classmethod
    def _resolve_file_data(cls, ctx: ExecutionContext, config: dict[str, Any]) -> dict[str, Any] | None:
        # 1) The explicit edge connected to the parser input is authoritative.
        direct = cls._candidate_to_file_data(ctx.get_input("document"))
        if direct:
            direct.setdefault("_selection", "connected_input")
            return direct

        # 2) Accept a transformed artifact connected under a non-standard target port.
        for input_name, input_value in ctx.inputs.items():
            if input_name.startswith("__") or input_name == "document":
                continue
            candidate = cls._candidate_to_file_data(
                input_value,
                provenance={"target_input": input_name, "distance": 0},
            )
            if candidate:
                candidate.setdefault("_selection", "connected_input_alias")
                return candidate

        # 3) Search all successful upstream outputs, ordered by graph distance. This
        # makes a parser after converter/cleaner nodes use the nearest (latest) file.
        lineage = ctx.get_input("__upstream_artifacts__", [])
        if isinstance(lineage, list):
            for artifact in lineage:
                if not isinstance(artifact, dict):
                    continue
                provenance = {
                    "source_node_id": artifact.get("source_node_id"),
                    "source_node_type": artifact.get("source_node_type"),
                    "source_output_port": artifact.get("output_port"),
                    "distance": artifact.get("distance"),
                }
                candidate = cls._candidate_to_file_data(
                    artifact.get("value"),
                    provenance=provenance,
                )
                if candidate:
                    candidate.setdefault("_selection", "nearest_upstream_artifact")
                    return candidate

        # 4) Standalone node test / manually configured fallback.
        file_id = str(config.get("file_id", "")).strip()
        if not file_id:
            return None

        upload_dir = cls._get_upload_dir()
        if not upload_dir.exists():
            return None

        index_path = upload_dir / ".documents-index.json"
        if index_path.exists():
            try:
                index = json.loads(index_path.read_text(encoding="utf-8"))
                record = index.get(file_id) if isinstance(index, dict) else None
                if isinstance(record, dict):
                    candidate = cls._candidate_to_file_data(record)
                    if candidate:
                        candidate.setdefault("file_id", file_id)
                        candidate.setdefault("_selection", "configured_file_id")
                        return candidate
            except (OSError, json.JSONDecodeError):
                pass

        for file_path in upload_dir.rglob("*"):
            if file_path.is_file() and file_path.stem == file_id:
                return {
                    "file_id": file_id,
                    "filename": str(config.get("filename") or file_path.name),
                    "path": str(file_path),
                    "file_path": str(file_path),
                    "size_bytes": file_path.stat().st_size,
                    "_selection": "configured_file_id",
                }
        return None

    @staticmethod
    def _read_text_file(path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "utf-16", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except (UnicodeDecodeError, UnicodeError):
                continue
        return path.read_text(errors="replace")

    @staticmethod
    def _parse_pdf(path: Path, preserve_page_breaks: bool) -> tuple[str, dict[str, Any]]:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages: list[dict[str, Any]] = []
        text_parts: list[str] = []
        empty_pages: list[int] = []

        for index, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            pages.append({
                "page": index,
                "text": page_text,
                "character_count": len(page_text),
                "empty": not bool(page_text.strip()),
            })
            if not page_text.strip():
                empty_pages.append(index)
            if preserve_page_breaks:
                text_parts.append(f"--- Page {index} ---\n{page_text}".rstrip())
            else:
                text_parts.append(page_text)

        separator = "\n\n" if preserve_page_breaks else "\n"
        return separator.join(text_parts), {
            "pages": pages,
            "empty_pages": empty_pages,
            "requires_ocr": bool(empty_pages),
        }

    @staticmethod
    def _iter_docx_blocks(document: Any) -> Iterable[tuple[str, Any]]:
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                yield "paragraph", Paragraph(child, document)
            elif tag == "tbl":
                yield "table", Table(child, document)

    @classmethod
    def _parse_docx(
        cls,
        path: Path,
        include_tables: bool,
        include_headers_footers: bool,
    ) -> tuple[str, dict[str, Any]]:
        import docx

        document = docx.Document(str(path))
        paragraphs: list[str] = []
        tables: list[dict[str, Any]] = []
        blocks: list[dict[str, Any]] = []
        text_parts: list[str] = []
        paragraph_index = 0
        table_index = 0

        for block_type, block in cls._iter_docx_blocks(document):
            if block_type == "paragraph":
                paragraph_index += 1
                text = block.text or ""
                paragraphs.append(text)
                blocks.append({
                    "type": "paragraph",
                    "index": paragraph_index,
                    "text": text,
                    "style": getattr(getattr(block, "style", None), "name", None),
                })
                if text:
                    text_parts.append(text)
                continue

            table_index += 1
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            table_payload = {"table": table_index, "rows": rows}
            if include_tables:
                tables.append(table_payload)
                blocks.append({"type": "table", **table_payload})
                text_parts.extend("\t".join(row) for row in rows)
            else:
                blocks.append({"type": "table", "table": table_index, "row_count": len(rows), "included": False})

        headers: list[dict[str, Any]] = []
        footers: list[dict[str, Any]] = []
        if include_headers_footers:
            seen_header_parts: set[str] = set()
            seen_footer_parts: set[str] = set()
            for section_index, section in enumerate(document.sections, start=1):
                header_key = str(section.header.part.partname)
                if header_key not in seen_header_parts:
                    seen_header_parts.add(header_key)
                    header_text = "\n".join(p.text for p in section.header.paragraphs if p.text)
                    header_tables = [
                        [[cell.text for cell in row.cells] for row in table.rows]
                        for table in section.header.tables
                    ]
                    headers.append({"section": section_index, "text": header_text, "tables": header_tables})
                    if header_text:
                        text_parts.insert(0, f"[Header {section_index}]\n{header_text}")
                    if include_tables:
                        for rows in reversed(header_tables):
                            text_parts.insert(0, "\n".join("\t".join(row) for row in rows))

                footer_key = str(section.footer.part.partname)
                if footer_key not in seen_footer_parts:
                    seen_footer_parts.add(footer_key)
                    footer_text = "\n".join(p.text for p in section.footer.paragraphs if p.text)
                    footer_tables = [
                        [[cell.text for cell in row.cells] for row in table.rows]
                        for table in section.footer.tables
                    ]
                    footers.append({"section": section_index, "text": footer_text, "tables": footer_tables})
                    if footer_text:
                        text_parts.append(f"[Footer {section_index}]\n{footer_text}")
                    if include_tables:
                        text_parts.extend("\n".join("\t".join(row) for row in rows) for rows in footer_tables)

        text = "\n".join(part for part in text_parts if part is not None)
        return text, {
            "blocks": blocks,
            "paragraphs": paragraphs,
            "tables": tables,
            "headers": headers,
            "footers": footers,
        }

    @staticmethod
    def _normalise_cell(value: Any) -> Any:
        if value is None:
            return ""
        if hasattr(value, "isoformat"):
            try:
                return value.isoformat()
            except (TypeError, ValueError):
                pass
        return value

    @classmethod
    def _parse_xlsx(
        cls,
        path: Path,
        include_tables: bool,
        sheet_name: str,
        max_rows: int | None,
    ) -> tuple[str, dict[str, Any]]:
        from openpyxl import load_workbook

        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        selected_names = [sheet_name] if sheet_name else list(workbook.sheetnames)
        missing = [name for name in selected_names if name not in workbook.sheetnames]
        if missing:
            workbook.close()
            raise ValueError(f"Excel sheet not found: {missing[0]}")

        sheets: list[dict[str, Any]] = []
        text_parts: list[str] = []
        any_truncated = False

        for name in selected_names:
            worksheet = workbook[name]
            rows: list[list[Any]] = []
            truncated = False
            for index, raw_row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if max_rows is not None and index > max_rows:
                    truncated = True
                    break
                row = [cls._normalise_cell(value) for value in raw_row]
                while row and row[-1] == "":
                    row.pop()
                # Preserve non-empty rows; completely blank rows carry no parseable data.
                if row:
                    rows.append(row)

            any_truncated = any_truncated or truncated
            text_parts.append(f"[{name}]")
            text_parts.extend("\t".join(str(value) for value in row) for row in rows)
            sheets.append(
                {
                    "name": name,
                    "rows": rows if include_tables else [],
                    "row_count": len(rows),
                    "truncated": truncated,
                }
            )

        workbook.close()
        return "\n".join(text_parts), {"sheets": sheets, "truncated": any_truncated}

    @classmethod
    def _parse_csv(
        cls,
        path: Path,
        include_tables: bool,
        max_rows: int | None,
    ) -> tuple[str, dict[str, Any]]:
        content = cls._read_text_file(path)
        sample = content[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel

        rows: list[list[str]] = []
        truncated = False
        reader = csv.reader(io.StringIO(content), dialect=dialect)
        for index, row in enumerate(reader, start=1):
            if max_rows is not None and index > max_rows:
                truncated = True
                break
            rows.append(row)

        text = "\n".join("\t".join(row) for row in rows)
        return text, {
            "rows": rows if include_tables else [],
            "row_count": len(rows),
            "truncated": truncated,
        }

    @staticmethod
    def _safe_stem(filename: str) -> str:
        stem = Path(filename).stem
        cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-._")
        return cleaned or "document"

    @classmethod
    def _create_download_outputs(
        cls,
        *,
        source_filename: str,
        text: str,
        data: dict[str, Any],
        metadata: dict[str, Any],
    ) -> dict[str, dict[str, Any]]:
        output_root = cls._get_storage_dir() / "parser_outputs"
        output_root.mkdir(parents=True, exist_ok=True)

        token = uuid.uuid4().hex
        stem = cls._safe_stem(source_filename)
        text_name = f"{stem}-parsed-{token[:8]}.txt"
        json_name = f"{stem}-structured-{token[:8]}.json"
        zip_name = f"{stem}-parser-output-{token[:8]}.zip"
        text_path = output_root / text_name
        json_path = output_root / json_name
        zip_path = output_root / zip_name

        text_path.write_text(text, encoding="utf-8")
        json_payload = {
            "metadata": metadata,
            "text": text,
            "data": data,
        }
        json_path.write_text(
            json.dumps(json_payload, ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )

        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.write(text_path, arcname=text_name)
            archive.write(json_path, arcname=json_name)

        def descriptor(path: Path, mime_type: str) -> dict[str, Any]:
            relative = path.relative_to(cls._get_storage_dir()).as_posix()
            return {
                "filename": path.name,
                "path": str(path),
                "download_url": f"/data/storage/{relative}",
                "mime_type": mime_type,
                "size_bytes": path.stat().st_size,
            }

        return {
            "text": descriptor(text_path, "text/plain; charset=utf-8"),
            "json": descriptor(json_path, "application/json"),
            "bundle": descriptor(zip_path, "application/zip"),
        }

    async def execute(self, ctx: ExecutionContext) -> NodeResult:
        result = NodeResult(node_id=self.node_id)
        result.start()

        try:
            config = self.get_resolved_config()
            file_data = self._resolve_file_data(ctx, config)
            if not file_data:
                result.fail(
                    "No document found. Connect a Document Input or document-processing node, "
                    "or select a file for a standalone test."
                )
                return result

            path_value = file_data.get("path") or file_data.get("file_path")
            if not path_value:
                result.fail("No file path found in the latest upstream document")
                return result

            file_path = Path(str(path_value))
            if not file_path.exists():
                result.fail(f"File not found at path: {file_path}")
                return result

            extension = file_path.suffix.lower()
            supported = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md"}
            if extension not in supported:
                result.fail(
                    f"File extension '{extension}' is not supported by the parser. Supported: {sorted(supported)}. "
                    "Images must first go through OCR, and legacy .doc/.xls files must be converted."
                )
                return result

            max_bytes = int(config.get("max_file_size_mb", 50)) * 1024 * 1024
            size_bytes = int(file_data.get("size_bytes", file_path.stat().st_size))
            if size_bytes > max_bytes:
                result.fail(f"File size {size_bytes} bytes exceeds maximum of {max_bytes} bytes")
                return result

            parse_entire = bool(config.get("parse_entire_document", True))
            include_tables = bool(config.get("include_tables", True))
            include_headers_footers = bool(config.get("include_headers_footers", True))
            preserve_page_breaks = bool(config.get("preserve_page_breaks", True))
            configured_max_rows = max(1, int(config.get("max_rows_per_sheet", 5000)))
            max_rows = None if parse_entire else configured_max_rows
            sheet_name = str(config.get("sheet_name", "")).strip()

            if extension == ".pdf":
                text, data = self._parse_pdf(file_path, preserve_page_breaks)
                engine = "pypdf"
            elif extension == ".docx":
                text, data = self._parse_docx(file_path, include_tables, include_headers_footers)
                engine = "python-docx"
            elif extension in {".xlsx", ".xlsm"}:
                text, data = self._parse_xlsx(file_path, include_tables, sheet_name, max_rows)
                engine = "openpyxl"
            elif extension == ".csv":
                text, data = self._parse_csv(file_path, include_tables, max_rows)
                engine = "python-csv"
            else:
                text = self._read_text_file(file_path)
                data = {"lines": text.splitlines(), "truncated": False}
                engine = "python-text"

            original_name = str(file_data.get("filename") or file_path.name)
            provenance = file_data.get("_provenance") if isinstance(file_data.get("_provenance"), dict) else {}
            truncated = bool(data.get("truncated"))
            if "sheets" in data:
                truncated = truncated or any(bool(sheet.get("truncated")) for sheet in data["sheets"])

            metadata: dict[str, Any] = {
                "filename": original_name,
                "source_path": str(file_path),
                "extension": extension,
                "parser_engine": engine,
                "size_bytes": size_bytes,
                "character_count": len(text),
                "line_count": len(text.splitlines()),
                "parse_entire_document": parse_entire,
                "complete": not truncated,
                "truncated": truncated,
                "include_tables": include_tables,
                "selected_source": file_data.get("_selection", "connected_input"),
                **{key: value for key, value in provenance.items() if value is not None},
            }

            if "pages" in data:
                metadata["page_count"] = len(data["pages"])
                metadata["empty_page_count"] = len(data.get("empty_pages", []))
                metadata["requires_ocr"] = bool(data.get("requires_ocr"))
            if "sheets" in data:
                metadata["sheet_count"] = len(data["sheets"])
                metadata["row_count"] = sum(int(sheet.get("row_count", 0)) for sheet in data["sheets"])
            if "tables" in data:
                metadata["table_count"] = len(data["tables"])
            if "blocks" in data:
                metadata["block_count"] = len(data["blocks"])
            if "rows" in data:
                metadata["row_count"] = int(data.get("row_count", 0))

            downloads: dict[str, dict[str, Any]] = {}
            if bool(config.get("create_downloads", True)):
                downloads = self._create_download_outputs(
                    source_filename=original_name,
                    text=text,
                    data=data,
                    metadata=metadata,
                )

            result.succeed({
                "text": text,
                "data": data,
                "metadata": metadata,
                "downloads": downloads,
            })
        except ImportError as exc:
            result.fail(
                f"Document parser dependency missing: {exc}. Install pypdf, python-docx, and openpyxl."
            )
        except Exception as exc:
            result.fail(str(exc))

        return result
